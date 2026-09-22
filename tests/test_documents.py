"""Extraction of source documents the `/wiki/raw/` tier accepts.

The tier holds what people put there, and people put Word files there. These
cover the extraction itself; `test_volume_backend.py` covers the two places the
backend calls it.
"""

from __future__ import annotations

import io

import pytest

from agent_server.documents import (
    UnreadableDocumentError,
    decode,
    docx_text,
    is_extractable,
)


def _docx(build) -> bytes:
    """A real `.docx` in memory, built by the same library that reads it."""
    from docx import Document

    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── format detection ─────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("path", "expected"),
    [
        ("/policies/sop.docx", True),
        ("/policies/SOP.DOCX", True),  # the Volume preserves case; we must not care
        ("/policies/sop.md", False),
        ("/policies/sop.docx.md", False),  # suffix, not substring
    ],
)
def test_only_word_documents_are_extracted(path, expected):
    assert is_extractable(path) is expected


# ── extraction ───────────────────────────────────────────────────────────────


def test_paragraphs_come_back_as_text():
    raw = _docx(lambda d: [d.add_paragraph("Tujuan"), d.add_paragraph("Ruang Lingkup")])
    text = docx_text(raw)
    assert "Tujuan" in text and "Ruang Lingkup" in text


def test_a_table_is_rendered_as_markdown():
    """A policy document keeps its targets in a table, so the rows are the part
    that must survive — and markdown is how every other document here states
    them."""

    def build(d):
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Prioritas"
        table.cell(0, 1).text = "Target"
        table.cell(1, 0).text = "P1"
        table.cell(1, 1).text = "4 jam"

    text = docx_text(_docx(build))
    assert "| Prioritas | Target |" in text
    assert "|---|---|" in text
    assert "| P1 | 4 jam |" in text


def test_blocks_keep_document_order():
    """python-docx exposes paragraphs and tables as separate collections, so
    reading them in turn would move every table to the end — detaching a target
    table from the heading that says what it targets."""

    def build(d):
        d.add_paragraph("SEBELUM")
        d.add_table(rows=1, cols=1).cell(0, 0).text = "TENGAH"
        d.add_paragraph("SESUDAH")

    text = docx_text(_docx(build))
    assert text.index("SEBELUM") < text.index("TENGAH") < text.index("SESUDAH")


def test_a_cell_containing_a_pipe_does_not_break_the_table():
    def build(d):
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "a | b"

    assert r"a \| b" in docx_text(_docx(build))


def test_an_empty_document_is_empty_not_an_error():
    assert docx_text(_docx(lambda d: None)) == ""


# ── failure is a message, not an exception ───────────────────────────────────


def test_a_file_that_is_not_a_word_document_is_refused_not_raised():
    """Raised as our own type so the backend reports it as a tool result the
    agent can act on, rather than as a 500 from the route."""
    with pytest.raises(UnreadableDocumentError):
        docx_text(b"this is not a zip", "/policies/broken.docx")


def test_the_non_utf8_message_still_names_utf8():
    with pytest.raises(UnreadableDocumentError, match="UTF-8"):
        decode(b"\xff\xfe\x00binary", "/blob.md")


def test_markdown_still_decodes_unchanged():
    assert decode(b"# Judul\n", "/index.md") == "# Judul\n"
